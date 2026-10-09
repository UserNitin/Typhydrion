"""Rename the project in the report from Typhydrion to an IEEE-style descriptive title."""
from __future__ import annotations

import shutil
import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageEnhance

sys.path.insert(0, str(Path(__file__).parent))
import insert_report_figures as figures  # noqa: E402

REPORT = Path(r"C:\Users\NITIN\Downloads\Typhydrion_Visual_ML_Pipeline_Builder_Report_Through_Preprocessing.docx")
BACKUP = REPORT.with_name("Typhydrion_Report_Before_Rename.docx")

OLD_NAME = "Typhydrion"
NEW_NAME = figures.PROJECT_NAME
FULL_TITLE = f"{NEW_NAME}: A Node-Based Visual Framework for Building Machine Learning Data Preprocessing Pipelines"
OLD_TITLES = (
    "Typhydrion \u2013 Visual Node-Based ML Pipeline Builder",
    "Typhydrion \u2014 Visual Node-Based ML Pipeline Builder",
    "Typhydrion - Visual Node-Based ML Pipeline Builder",
)


def all_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs


def rename_text(doc: Document) -> int:
    count = 0
    for paragraph in all_paragraphs(doc):
        for run in paragraph.runs:
            text = run.text
            for old in OLD_TITLES:
                text = text.replace(old, FULL_TITLE)
            text = text.replace(OLD_NAME, NEW_NAME)
            if text != run.text:
                count += 1
                run.text = text
    return count


def remove_spacer_after_title(doc: Document) -> None:
    """The longer title wraps to two lines; drop one blank line so the page still fits."""
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.startswith("Project Title:"):
            following = paragraphs[index + 1 : index + 3]
            if len(following) == 2 and all(not p.text.strip() and "graphic" not in p._element.xml for p in following):
                following[0]._element.getparent().remove(following[0]._element)
            return


def png_bytes(path: Path, grayscale: bool) -> bytes:
    image = Image.open(path)
    if grayscale:
        image = ImageEnhance.Contrast(image.convert("L")).enhance(1.2)
    buffer = BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def replace_figures(doc: Document) -> None:
    images = {
        "Figure 1:": png_bytes(figures.make_architecture(), grayscale=False),
        "Figure 2:": png_bytes(figures.make_dfd_level0(), grayscale=True),
        "Figure 3:": png_bytes(figures.make_dfd_level1(), grayscale=True),
    }
    paragraphs = doc.paragraphs
    replaced = 0
    for index, paragraph in enumerate(paragraphs):
        prefix = next((p for p in images if paragraph.text.strip().startswith(p)), None)
        if prefix is None:
            continue
        for candidate in paragraphs[index + 1 : index + 4]:
            blips = candidate._element.findall(".//" + qn("a:blip"))
            if blips:
                doc.part.related_parts[blips[0].get(qn("r:embed"))]._blob = images[prefix]
                replaced += 1
                break
    if replaced != 3:
        raise RuntimeError(f"Expected to replace 3 figures, replaced {replaced}")


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else REPORT
    doc = Document(str(REPORT))
    print("Runs renamed:", rename_text(doc))
    remove_spacer_after_title(doc)
    replace_figures(doc)
    doc.core_properties.title = FULL_TITLE
    if out == REPORT and not BACKUP.exists():
        shutil.copy2(REPORT, BACKUP)
    doc.save(str(out))
    print("Saved:", out)


if __name__ == "__main__":
    main()
