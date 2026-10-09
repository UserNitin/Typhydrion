"""Expand the Typhydrion abstract to fill its dedicated report page."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


DOWNLOADS = Path(r"C:\Users\NITIN\Downloads")
PROJECT = Path(r"d:\collage\major project")
SOURCE = DOWNLOADS / "Typhydrion_Visual_ML_Pipeline_Builder_Report_Expanded.docx"
OUTPUT = DOWNLOADS / "Typhydrion_Visual_ML_Pipeline_Builder_Report_Expanded_Abstract.docx"
PROJECT_OUTPUT = PROJECT / OUTPUT.name


def format_run(run) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")


def insert_before(anchor: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._element.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = "Normal"
    run = paragraph.add_run(text)
    format_run(run)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_after = Pt(6)
    return paragraph


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    shutil.copy2(SOURCE, PROJECT_OUTPUT)
    doc = Document(str(PROJECT_OUTPUT))
    marker = "The internal execution design separates the visual graph"
    if any(p.text.strip().startswith(marker) for p in doc.paragraphs):
        print("Abstract already expanded")
    else:
        introduction = next(
            (p for p in doc.paragraphs if p.text.strip() == "2. INTRODUCTION"),
            None,
        )
        if introduction is None:
            raise RuntimeError("Could not find the Introduction insertion point")

        additions = [
            (
                "The internal execution design separates the visual graph from node runtime "
                "logic. Graph validation checks structural requirements and rejects cyclic or "
                "incomplete connections. The scheduler derives dependency levels, while the "
                "pipeline executor and node runner pass upstream outputs into registered runtime "
                "classes. Each runtime follows a common result contract containing outputs, "
                "metadata, execution status, timing information, and an error message when a "
                "step cannot complete. This separation allows the editor to remain interactive "
                "while preprocessing behavior is organized into reusable components."
            ),
            (
                "A representative workflow begins with the Dataset Loader, which reads a "
                "user-selected file and emits raw data, schema information, aggregate statistics, "
                "and feature or target candidates. Inspection windows display rows, data types, "
                "missing percentages, distributions, and profiling information. Downstream nodes "
                "then apply declared operations such as median or mode imputation, type coercion, "
                "one-hot encoding, standardization, outlier treatment, and feature selection. "
                "Intermediate outputs remain accessible through node ports and output windows, "
                "allowing the operator to compare the data before and after each transformation."
            ),
            (
                "The implemented phase was demonstrated with a small mixed-type tabular dataset "
                "containing numeric, categorical, and missing values. The loader produced a "
                "twelve-row, six-column frame and reported its quality characteristics. A "
                "missing-value step generated per-column counts and percentages, categorical "
                "fields were converted through one-hot encoding, and numeric fields were scaled. "
                "The resulting prepared frame contained eleven columns after encoding. These "
                "observations are used as functional evidence of data movement and transformation; "
                "they are not presented as predictive accuracy results."
            ),
            (
                "The principal outcome is an inspectable preprocessing environment in which "
                "configuration, execution order, transformed data, and supporting metadata can be "
                "reviewed together. The approach is intended for academic experimentation and "
                "local analytical work where reproducibility and visibility are more important "
                "than automated model selection. Future phases can extend the same graph with "
                "split, training, evaluation, export, and inference nodes, provided that fitted "
                "preprocessing state is consistently reused and validated on holdout data."
            ),
        ]
        for text in additions:
            insert_before(introduction, text)

    doc.save(str(PROJECT_OUTPUT))
    shutil.copy2(PROJECT_OUTPUT, OUTPUT)
    print("Saved:", OUTPUT)


if __name__ == "__main__":
    main()
