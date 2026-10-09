"""Expand Sections 6 and 8 of the Typhydrion report so their last pages are filled."""
from __future__ import annotations

import copy
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

REPORT = Path(r"C:\Users\NITIN\Downloads\Typhydrion_Visual_ML_Pipeline_Builder_Report_Through_Preprocessing.docx")
BACKUP = REPORT.with_name("Typhydrion_Report_Before_Sections_6_8_Expansion.docx")

SECTION_6 = [
    ("h", "6.6 Demonstration dataset"),
    ("p", "The screenshots in Section 9 use a small demonstration file, sample_preprocess.csv, "
          "with 12 rows and six columns: age, income, city, plan, score and visits. Four columns "
          "are numeric and two are categorical. Ten cells are deliberately left empty across all six "
          "columns so that the Missing Value Handler and its Missing Report have visible work to do, "
          "and the categorical columns repeat labels such as Bengaluru, Delhi, Mumbai, Chennai, Basic, "
          "Plus and Pro so that encoding produces meaningful indicator columns."),
    ("p", "After median imputation, one-hot encoding of city and plan, and StandardScaler on the numeric "
          "columns, the frame grows from 6 to 11 columns while keeping all 12 rows, because the pipeline "
          "fills incomplete records instead of dropping them. The small size lets every value be checked "
          "by hand against the preview windows."),
    ("p", "Dataset handling in this phase assumes that the file fits in memory, that the first row holds "
          "column names, and that each loader node reads one table. Very large files, workbooks whose "
          "sheets use different headers, and free-text columns that need tokenization are only partly "
          "supported, through reader keyword arguments and light text normalization."),
]

SECTION_8_1 = [
    ("p", "Read from left to right, the Level-1 diagram follows the order in which the executor runs the "
          "nodes. The loader turns the source file into a DataFrame and records its schema. The preview "
          "and profile step only reads that frame to fill the Data Preview, Data Statistics and profiling "
          "windows. The cleaning step handles missing values, data types and outliers; encoding and "
          "scaling then convert the cleaned columns into numeric features and keep their fitted state on "
          "the node; and feature selection optionally narrows the result before it reaches the output ports."),
    ("p", "Each arrow corresponds to a port connection in the node editor, and the data carried on it is a "
          "pandas DataFrame together with a small metadata dictionary. Because nodes run in topological "
          "order, a downstream node never receives a partial result: it gets either the complete output of "
          "its upstream node or an error."),
    ("p", "The diagram also shows where the user's choices enter the flow. Settings from the Node Properties "
          "window, such as the imputation strategy, the encoding method or the scaler type, are read by each "
          "node at run time and are not stored in the data itself. Changing a setting and running the graph "
          "again therefore repeats the whole flow from the source file, so every output can be traced to a "
          "specific configuration."),
]

SECTION_8_3 = [
    ("h", "8.3 Data flow rules and error handling"),
    ("p", "Three rules govern every flow in the diagram. First, data moves only along connected ports, so "
          "nodes do not share hidden state. Second, every node returns a NodeResult that carries either "
          "output data or an error message, and the executor stops the dependent branch when an error "
          "occurs. Third, inspection windows read node outputs but never write back to them, which keeps "
          "viewing separate from transformation."),
    ("p", "Errors are reported to the user through the failing node's message and are also written to the "
          "logs folder with a timestamp for later review. Data leaves the system boundary only when the "
          "user asks for it, either through the Export node or as an HTML profiling report saved in the "
          "reports folder, so nothing is uploaded or shared automatically."),
    ("p", "Together these rules make the diagrams useful for testing as well as documentation. Each arrow can "
          "be checked by opening the node output window at that point and comparing row counts, column names "
          "and missing-value totals with the previous step. In the demonstration run, this check confirms that "
          "the row count stays at 12 throughout and that the column count rises from 6 to 11 after encoding. "
          "The same check can be repeated on any new dataset, so the diagrams remain a reliable guide when "
          "further preprocessing nodes are added to the pipeline in later phases of the project."),
]


def find(doc: Document, prefix: str) -> int:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i
    raise ValueError(f"Paragraph not found: {prefix}")


def clone(template: Paragraph, text: str):
    element = copy.deepcopy(template._element)
    runs = element.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r")
    for extra in runs[1:]:
        element.remove(extra)
    for child in list(element):
        if child.tag.endswith("}hyperlink") or child.tag.endswith("}bookmarkStart") or child.tag.endswith("}bookmarkEnd"):
            element.remove(child)
    paragraph = Paragraph(element, template._parent)
    paragraph.runs[0].text = text
    return element


def insert_after(anchor: Paragraph, items, heading_tpl: Paragraph, body_tpl: Paragraph) -> None:
    current = anchor._element
    for kind, text in items:
        element = clone(heading_tpl if kind == "h" else body_tpl, text)
        current.addnext(element)
        current = element


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else REPORT
    doc = Document(str(REPORT))
    if any(p.text.strip() == "6.6 Demonstration dataset" for p in doc.paragraphs):
        print("Already expanded")
        return

    ps = doc.paragraphs
    heading_tpl = ps[find(doc, "6.5 Transformations")]
    body_tpl = ps[find(doc, "Chronological or random train/test")]

    insert_after(ps[find(doc, "D1 is the source file on disk")], SECTION_8_3, heading_tpl, body_tpl)
    insert_after(ps[find(doc, "At Level 1, the graph is validated")], SECTION_8_1, heading_tpl, body_tpl)
    insert_after(body_tpl, SECTION_6, heading_tpl, body_tpl)

    if out == REPORT and not BACKUP.exists():
        shutil.copy2(REPORT, BACKUP)
    doc.save(str(out))
    print("Saved:", out)


if __name__ == "__main__":
    main()
