"""Expand Sections 3 and 7, apply 13 pt text, and grayscale DFD figures."""
from __future__ import annotations

import io
import shutil
from pathlib import Path

from PIL import Image, ImageEnhance
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


DOWNLOADS = Path(r"C:\Users\NITIN\Downloads")
PROJECT = Path(r"d:\collage\major project")
SOURCE = DOWNLOADS / "Typhydrion_Visual_ML_Pipeline_Builder_Report_Expanded_Abstract.docx"
OUTPUT = DOWNLOADS / "Typhydrion_Report_Sections_3_7_Expanded_13pt_Grayscale_DFD.docx"
PROJECT_OUTPUT = PROJECT / OUTPUT.name


SECTION_3 = [
    (
        "3.3 Root causes and problem decomposition",
        "Heading 2",
    ),
    (
        "The problem can be divided into four connected causes. The first is data-quality "
        "variation: files may contain null values, mixed types, duplicate categories, extreme "
        "numeric observations, or columns that are irrelevant to a later model. The second is "
        "process fragmentation, where loading, cleaning, and transformation are performed in "
        "different tools without a single record of their order. The third is weak observability: "
        "a transformed table may be produced without a report explaining which values were "
        "changed. The fourth is execution risk, because an invalid dependency or cyclic workflow "
        "can prevent a pipeline from producing a reliable terminal output. Typhydrion addresses "
        "these causes through dedicated data and preprocessing nodes, intermediate output "
        "inspection, and validation before execution. A missing path, empty required port, or "
        "cyclic dependency is therefore reported at the component responsible for it.",
        "Normal",
    ),
]


SECTION_7 = [
    (
        "7.3 Integration of the technology stack",
        "Heading 2",
    ),
    (
        "The technologies are connected through a layered implementation rather than used as "
        "isolated libraries. PySide6 provides the main window, node editor, dialogs, property "
        "controls, data tables, statistics panels, and output views. User actions in this layer "
        "produce graph and node configurations. The execution layer then resolves the selected "
        "visual node through the registry and invokes the corresponding Python runtime class. "
        "This arrangement keeps interface concerns separate from transformation logic while "
        "allowing runtime results to be returned to the visual windows.",
        "Normal",
    ),
]


def set_run_font(run, size: float = 13, bold: bool | None = None) -> None:
    if bold is not None:
        run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{key}"), "Times New Roman")


def insert_before(anchor: Paragraph, text: str, style: str) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._element.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    run = paragraph.add_run(text)
    set_run_font(run, 13, bold=style.startswith("Heading"))
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT if style.startswith("Heading") else WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(6 if style.startswith("Heading") else 0)
    fmt.space_after = Pt(6)
    return paragraph


def insert_section_content(doc: Document, marker: str, next_heading: str, additions: list) -> None:
    if any(p.text.strip() == marker for p in doc.paragraphs):
        return
    anchor = next((p for p in doc.paragraphs if p.text.strip() == next_heading), None)
    if anchor is None:
        raise RuntimeError(f"Could not find insertion point: {next_heading}")
    for text, style in additions:
        insert_before(anchor, text, style)


def apply_thirteen_point(doc: Document) -> None:
    """Use 13 pt in Sections 3 and 7 while preserving the established report layout."""
    active = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in ("3. PROBLEM STATEMENT", "7. TECHNOLOGIES USED"):
            active = True
        elif text in ("4. EXISTING SYSTEM", "8. DATA FLOW DIAGRAM"):
            active = False
        if active or text.startswith("Figure 2:") or text.startswith("Figure 3:"):
            style = paragraph.style.name if paragraph.style else ""
            size = 14 if style == "Heading 1" else 13
            for run in paragraph.runs:
                set_run_font(run, size)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    # Table 6 belongs to Section 7 and follows the same 13 pt requirement.
    caption = next(
        (
            p
            for p in doc.paragraphs
            if p.text.strip().startswith("Table 6: Technologies used")
        ),
        None,
    )
    if caption is not None:
        sibling = caption._element.getnext()
        while sibling is not None and sibling.tag != qn("w:tbl"):
            sibling = sibling.getnext()
        if sibling is not None:
            table = next((t for t in doc.tables if t._tbl is sibling), None)
            if table is not None:
                for row_index, row in enumerate(table.rows):
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                set_run_font(
                                    run,
                                    13,
                                    bold=True if row_index == 0 else None,
                                )


def grayscale_image_part(part) -> None:
    image = Image.open(io.BytesIO(part.blob))
    # Convert to high-contrast grayscale while retaining anti-aliased text.
    gray = ImageEnhance.Contrast(image.convert("L")).enhance(1.2)
    output = io.BytesIO()
    gray.save(output, format="PNG", optimize=True)
    part._blob = output.getvalue()
    part._content_type = "image/png"


def grayscale_dfd_figures(doc: Document) -> None:
    """Replace pictures immediately following Figure 2 and Figure 3 captions."""
    paragraphs = doc.paragraphs
    changed = 0
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not (text.startswith("Figure 2:") or text.startswith("Figure 3:")):
            continue
        for candidate in paragraphs[index + 1 : index + 4]:
            blips = candidate._element.findall(".//" + qn("a:blip"))
            if not blips:
                if candidate.text.strip():
                    break
                continue
            rel_id = blips[0].get(qn("r:embed"))
            part = doc.part.related_parts[rel_id]
            grayscale_image_part(part)
            changed += 1
            break
    if changed != 2:
        raise RuntimeError(f"Expected to grayscale 2 DFD figures, changed {changed}")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    shutil.copy2(SOURCE, PROJECT_OUTPUT)
    doc = Document(str(PROJECT_OUTPUT))
    insert_section_content(
        doc,
        "3.3 Root causes and problem decomposition",
        "4. EXISTING SYSTEM",
        SECTION_3,
    )
    insert_section_content(
        doc,
        "7.3 Integration of the technology stack",
        "8. DATA FLOW DIAGRAM",
        SECTION_7,
    )
    apply_thirteen_point(doc)
    grayscale_dfd_figures(doc)
    doc.save(str(PROJECT_OUTPUT))
    shutil.copy2(PROJECT_OUTPUT, OUTPUT)
    print("Saved:", OUTPUT)


if __name__ == "__main__":
    main()
