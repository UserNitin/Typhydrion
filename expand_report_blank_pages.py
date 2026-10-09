"""Expand the Introduction so its nearly blank continuation page is properly filled."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


DOWNLOADS = Path(r"C:\Users\NITIN\Downloads")
PROJECT = Path(r"d:\collage\major project")
NAME = "Typhydrion_Visual_ML_Pipeline_Builder_Report_Through_Preprocessing.docx"
SOURCE = DOWNLOADS / NAME
PROJECT_COPY = PROJECT / NAME
BACKUP = DOWNLOADS / "Typhydrion_Report_Before_Page_Expansion.docx"


def set_font(run, *, bold: bool = False, size: float = 12) -> None:
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{key}"), "Times New Roman")


def insert_before(anchor: Paragraph, text: str, *, style: str = "Normal") -> Paragraph:
    element = OxmlElement("w:p")
    anchor._element.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    run = paragraph.add_run(text)
    set_font(run, bold=style.startswith("Heading"), size=12)
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT if style.startswith("Heading") else WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(6 if style.startswith("Heading") else 0)
    fmt.space_after = Pt(6)
    return paragraph


def expand(path: Path) -> None:
    doc = Document(str(path))
    if any(p.text.strip() == "2.4 Significance of the proposed work" for p in doc.paragraphs):
        print("Introduction already expanded:", path)
        return

    section_three = next(
        (p for p in doc.paragraphs if p.text.strip() == "3. PROBLEM STATEMENT"),
        None,
    )
    if section_three is None:
        raise RuntimeError("Could not locate Section 3 insertion point")

    additions = [
        (
            "2.4 Significance of the proposed work",
            "Heading 2",
        ),
        (
            "The significance of Typhydrion lies in making preprocessing decisions visible. "
            "In a conventional script, a reader may need to trace several functions to discover "
            "which columns were imputed, how categories were encoded, or whether scaling was "
            "applied before feature selection. In the node editor, every operation occupies an "
            "explicit position in the graph and exposes its configuration through the node card "
            "and properties window. This representation supports academic review because the "
            "workflow can be discussed step by step rather than treated as a single opaque program.",
            "Normal",
        ),
        (
            "The project also supports repeatability at the data-preparation stage. A saved graph "
            "preserves the intended order of operations, while the validator and scheduler derive "
            "an executable dependency order from the connections. Node outputs include not only "
            "transformed DataFrames but also evidence such as schemas, missing-value summaries, "
            "encoder mappings, scaler state, outlier masks, and feature scores. These artifacts help "
            "the user verify what changed and identify unsuitable choices before model training is "
            "introduced in a later phase.",
            "Normal",
        ),
        (
            "A local desktop design is useful for classroom demonstrations and exploratory work. "
            "The user can inspect tabular data, construct a pipeline, and review intermediate "
            "results without deploying a server or transferring a dataset to a remote service. "
            "This does not by itself guarantee privacy or production readiness, but it provides a "
            "clear boundary: source files, temporary data, logs, and profiler reports remain under "
            "the local project folders controlled by the operator.",
            "Normal",
        ),
        (
            "2.5 Design principles and expected outcomes",
            "Heading 2",
        ),
        (
            "The implementation follows four practical principles. First, composability allows a "
            "user to connect small nodes instead of relying on one fixed preprocessing routine. "
            "Second, inspectability requires important transformations to expose both their data "
            "output and supporting metadata. Third, controlled execution requires the graph to be "
            "validated as a directed acyclic graph and scheduled in dependency order. Fourth, "
            "graceful failure requires missing files, empty input ports, incompatible options, and "
            "other invalid conditions to return clear node errors instead of silently producing a "
            "partial result.",
            "Normal",
        ),
        (
            "For the phase documented in this report, the expected outcome is a prepared feature "
            "table whose transformation history can be understood from the graph. A successful "
            "demonstration should show that a tabular file is loaded, its shape and quality are "
            "inspected, missing values are treated, categorical values are converted into numerical "
            "representations, and numeric features are scaled or selected as required. The result is "
            "not yet a prediction; it is a consistent input that can be passed to split and modeling "
            "nodes in the next project phase.",
            "Normal",
        ),
        (
            "The expected academic benefit is therefore not merely automation. The system provides "
            "a visual record of preprocessing assumptions and makes intermediate evidence available "
            "for discussion, correction, and comparison. This encourages users to treat data quality "
            "and transformation order as deliberate engineering decisions rather than as hidden "
            "setup performed immediately before model training.",
            "Normal",
        ),
    ]

    for text, style in additions:
        insert_before(section_three, text, style=style)

    doc.save(str(path))
    print("Expanded:", path)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)
        print("Backup:", BACKUP)
    expanded_download = DOWNLOADS / (
        "Typhydrion_Visual_ML_Pipeline_Builder_Report_Expanded.docx"
    )
    # Work from the backup made immediately before expansion. This also works
    # when the original report is open and locked by Word.
    shutil.copy2(BACKUP, PROJECT_COPY)
    expand(PROJECT_COPY)
    shutil.copy2(PROJECT_COPY, expanded_download)
    print("Expanded Downloads copy:", expanded_download)


if __name__ == "__main__":
    main()
